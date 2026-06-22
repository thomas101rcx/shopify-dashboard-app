"""Shopify Dashboard — ETL + Overview."""

from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
import streamlit as st

import config as cfg
from merge import generate_updated_pivot_xlsx, merge_overview_with_etl, updated_overview_with_etl
from overview import (
    load_24y_25n,
    load_account_labels,
    load_count,
    load_duplicates,
    load_lost_accounts,
    load_overview,
)
from rules import load_rules
from transform import apply_rules

st.set_page_config(page_title="Shopify Dashboard", layout="wide")
st.title("Shopify Dashboard")

page = st.sidebar.radio("Page", ["ETL", "Overview", "Merge"], index=0)


def _read_excel(path: str, sheet: str | None = None) -> pd.DataFrame:
    last_err = None
    for engine in cfg.EXCEL_ENGINES:
        try:
            result = pd.read_excel(path, sheet_name=sheet, engine=engine)
            if isinstance(result, dict):
                if sheet and sheet in result:
                    return result[sheet]
                return next(iter(result.values()))
            return result
        except Exception as e:
            last_err = e
    raise last_err  # type: ignore[misc]


def _list_sheets(path: str) -> list[str]:
    try:
        import python_calamine as pc

        return pc.CalamineWorkbook.open_path(path).sheet_names  # type: ignore[attr-defined]
    except Exception:
        pass
    try:
        import openpyxl

        return openpyxl.load_workbook(path, read_only=True).sheetnames
    except Exception:
        return []


if page == "ETL":
    rules_file = st.sidebar.file_uploader(
        "Rules file (yaml / json / py)", type=["yaml", "yml", "json", "py"]
    )
    data_files = st.sidebar.file_uploader(
        "Data files (csv / xlsx)", type=["csv", "xlsx"], accept_multiple_files=True
    )

    rules = []
    if rules_file is not None:
        suffix = Path(rules_file.name).suffix.lower()
        tmp = Path(f"/tmp/rules_{rules_file.file_id}{suffix}")
        tmp.write_bytes(rules_file.getvalue())
        try:
            rules = load_rules(str(tmp))
            st.sidebar.success(f"Loaded {len(rules)} rule(s)")
        except Exception as e:
            st.sidebar.error(f"Rules error: {e}")

    if data_files:
        for f in data_files:
            suffix = Path(f.name).suffix.lower()
            tmp = Path(f"/tmp/{f.file_id}_{f.name}")
            tmp.write_bytes(f.getvalue())

            try:
                if suffix == ".csv":
                    df = pd.read_csv(tmp)
                else:
                    sheets = _list_sheets(str(tmp))
                    if len(sheets) > 1:
                        sheet = st.radio(
                            f"Sheet — {f.name}",
                            sheets,
                            horizontal=True,
                            key=f"sheet_{f.file_id}",
                        )
                    elif sheets:
                        sheet = sheets[0]
                    else:
                        sheet = None
                    df = _read_excel(str(tmp), sheet=sheet)
            except Exception as e:
                st.error(f"Failed to read {f.name}: {e}")
                continue

            st.subheader(f.name)
            st.caption(f"Raw shape: {df.shape[0]} rows × {df.shape[1]} cols")
            st.dataframe(df.head(50), use_container_width=True)

            if rules:
                try:
                    result = apply_rules(df.copy(), rules)
                except Exception as e:
                    st.error(f"Transform error: {e}")
                    continue
                st.caption(f"Result shape: {result.shape[0]} rows × {result.shape[1]} cols")
                st.dataframe(result, use_container_width=True)

                csv_bytes = result.to_csv(index=False).encode()
                buf = io.BytesIO()
                result.to_excel(buf, index=False)
                buf.seek(0)

                c1, c2, c3 = st.columns(3)
                c1.download_button(
                    "Download CSV", csv_bytes, f"{Path(f.name).stem}_out.csv", "text/csv"
                )
                c2.download_button(
                    "Download XLSX",
                    buf.read(),
                    f"{Path(f.name).stem}_out.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                if c3.button("Send to Merge", key=f"send_{f.file_id}"):
                    st.session_state["etl_result"] = result
                    st.success("Sent to Merge tab — switch to Merge tab to view")
            st.divider()
    else:
        st.info("Upload one or more CSV/XLSX files to begin.")

elif page == "Overview":
    st.subheader("Account Overview")
    st.markdown("Upload your current `pivot_table.xlsx` to view the dashboard.")

    pivot_file = st.file_uploader(
        "Pivot table file (xlsx)", type=["xlsx"], key="overview"
    )

    if pivot_file is None:
        st.info("Upload a pivot_table.xlsx file to view the dashboard.")
        st.stop()

    import tempfile
    suffix = Path(pivot_file.name).suffix.lower()
    tmp = Path(f"/tmp/overview_{pivot_file.file_id}{suffix}")
    tmp.write_bytes(pivot_file.getvalue())

    try:
        from openpyxl import load_workbook
        wb = load_workbook(tmp, data_only=True)
        sheets = wb.sheetnames
    except Exception as e:
        st.error(f"Failed to read file: {e}")
        st.stop()

    tab_overview, tab_lost, tab_count, tab_24y, tab_labels, tab_dupes = st.tabs(
        ["Overview", "Lost Account", "Count", "24Y 25N", "Labels", "Duplicates"]
    )

    with tab_overview:
        try:
            from overview import load_overview_styled, load_label_rules

            styled_rows = load_overview_styled(wb)
            label_rules = load_label_rules(wb)

            cols = cfg.OVERVIEW_COLS
            col_to_xlsx_idx = cfg.COL_TO_XLSX_IDX

            data_rows = []
            # Per-row, per-column style info: {(row_idx, col_name): style_dict}
            cell_styles = {}

            for row_idx, row in enumerate(styled_rows):
                vals = []
                for col_name in cols:
                    xlsx_idx = col_to_xlsx_idx[col_name]
                    cell = row.get(xlsx_idx, {})
                    vals.append(cell.get("value"))
                    if cell.get("fill") or cell.get("font_color") or cell.get("label_type"):
                        cell_styles[(row_idx, col_name)] = {
                            "fill": cell.get("fill"),
                            "font_color": cell.get("font_color"),
                            "label_type": cell.get("label_type"),
                        }
                data_rows.append(vals)

            df = pd.DataFrame(data_rows, columns=cols)

            for c in cfg.QTY_COLS:
                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).astype(int)

            df[cfg.COL_QTY_2026_TOTAL] = (
                df[cfg.COL_QTY_2026_Q1] + df[cfg.COL_QTY_2026_Q2]
                + df[cfg.COL_QTY_2026_Q3] + df[cfg.COL_QTY_2026_Q4]
            )

            def _growth(new, old):
                if old == 0:
                    return None
                return (new - old) / old

            growth_col_defs = [
                (cfg.COL_GROWTH_23_24, cfg.COL_QTY_2024, cfg.COL_QTY_2023),
                (cfg.COL_GROWTH_24_25, cfg.COL_QTY_2025, cfg.COL_QTY_2024),
                (cfg.COL_GROWTH_25_26, cfg.COL_QTY_2026_TOTAL, cfg.COL_QTY_2025),
            ]

            for gcol, new_col, old_col in growth_col_defs:
                computed = df.apply(lambda r: _growth(r[new_col], r[old_col]), axis=1)
                xlsx_idx = col_to_xlsx_idx[gcol]
                for i in range(len(df)):
                    original_styled = styled_rows[i].get(xlsx_idx, {})
                    original_val = original_styled.get("value")
                    is_numeric_original = original_styled.get("is_numeric", False)
                    if isinstance(original_val, str) and original_val.strip():
                        df.at[i, gcol] = original_val
                    elif is_numeric_original:
                        df.at[i, gcol] = original_val
                    else:
                        df.at[i, gcol] = computed.iloc[i]

            growth_col_set = set(cfg.GROWTH_COLS)

            def _fmt_cell(v, col_name=None):
                """Format cell: % for growth, int for qty, str for text labels."""
                if pd.isna(v) or v is None:
                    return ""
                if isinstance(v, str):
                    return v  # Text label (New_24, Lost, SOS, etc.)
                if isinstance(v, (int, float)):
                    if col_name in growth_col_set and v != 0:
                        return f"{v:.1%}"
                    return str(int(v))
                return str(v)

            # Build HTML table with styled cells
            def _style_cell(val, style, col_name=None):
                """Render cell value with color badge if styled."""
                if val is None or (isinstance(val, float) and pd.isna(val)):
                    val = ""
                label_type = (style or {}).get("label_type")
                fill = (style or {}).get("fill")
                font_color = (style or {}).get("font_color")

                badge_map = cfg.BADGE_STYLES

                if label_type in badge_map:
                    bg, fg, dot = badge_map[label_type]
                    display = _fmt_cell(val, col_name)
                    return f'<td style="background:{bg};color:{fg};font-weight:600;padding:4px 8px;text-align:center;">{dot} {display}</td>'

                if font_color and font_color not in ("#000000", "FF000000", None):
                    display = _fmt_cell(val, col_name)
                    return f'<td style="color:{font_color};font-weight:600;padding:4px 8px;">{display}</td>'

                if fill and fill not in (cfg.FILL_COLORS.get("FFF2CEEF"), None):
                    display = _fmt_cell(val, col_name)
                    return f'<td style="background:{fill};padding:4px 8px;">{display}</td>'

                display = _fmt_cell(val, col_name)
                # Right-align numbers, left-align text
                try:
                    float(display.replace("%", "").replace(",", ""))
                    return f'<td style="padding:4px 8px;text-align:right;">{display}</td>'
                except (ValueError, TypeError):
                    return f'<td style="padding:4px 8px;">{display}</td>'

            # Build HTML table
            html = '<table style="border-collapse:collapse;width:100%;font-size:0.9em;">'
            # Header
            html += '<tr style="background:#f0f0f0;font-weight:700;">'
            for col_name in cols:
                html += f'<th style="padding:6px 8px;border:1px solid #ddd;text-align:left;">{col_name}</th>'
            html += '</tr>'

            for i in range(len(df)):
                html += '<tr>'
                for col_name in cols:
                    val = df.at[i, col_name]
                    style = cell_styles.get((i, col_name))
                    html += _style_cell(val, style, col_name)
                html += '</tr>'
            html += '</table>'

            st.metric("Total accounts", len(df))

            # Count label types across all columns
            type_counts = {"new": 0, "lost": 0, "reactivated": 0}
            for style in cell_styles.values():
                lt = style.get("label_type")
                if lt in type_counts:
                    type_counts[lt] += 1
            if any(type_counts.values()):
                badges = []
                if type_counts["new"]:
                    badges.append(f"🟡 New: {type_counts['new']}")
                if type_counts["lost"]:
                    badges.append(f"🔴 Lost: {type_counts['lost']}")
                if type_counts["reactivated"]:
                    badges.append(f"🟢 Reactivated: {type_counts['reactivated']}")
                st.caption(" · ".join(badges))

            st.write(html, unsafe_allow_html=True)

            # Label rules from Account Label Definition
            if label_rules:
                with st.expander("Account Label Definitions"):
                    rule_data = []
                    for label_name, rule in label_rules.items():
                        rule_data.append({
                            "Label": label_name,
                            "Definition": rule.get("definition", ""),
                            "Note": rule.get("note", ""),
                        })
                    st.dataframe(pd.DataFrame(rule_data), use_container_width=True, hide_index=True)

        except Exception as e:
            st.error(f"Failed to load Overview: {e}")

    with tab_lost:
        try:
            ws = wb[cfg.SHEET_LOST_ACCOUNT]
            rows = []
            for r in range(3, ws.max_row + 1):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, cfg.OVERVIEW_NUM_COLS + 1)]
                if all(v is None for v in row_vals):
                    continue
                rows.append(row_vals)
            if rows:
                df = pd.DataFrame(rows, columns=cfg.OVERVIEW_COLS)
                st.dataframe(df, use_container_width=True, height=400)
        except Exception as e:
            st.error(f"Failed to load Lost Account: {e}")

    with tab_count:
        try:
            ws = wb[cfg.SHEET_COUNT]
            rows = []
            for r in range(2, ws.max_row + 1):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
                if all(v is None for v in row_vals):
                    continue
                rows.append(row_vals)
            if rows:
                df = pd.DataFrame(rows, columns=cfg.COUNT_COL_LABELS[: len(rows[0])])
                st.dataframe(df, use_container_width=True)
        except Exception as e:
            st.error(f"Failed to load Count: {e}")

    with tab_24y:
        try:
            ws = wb[cfg.SHEET_24Y_25N]
            rows = []
            for r in range(2, ws.max_row + 1):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
                if all(v is None for v in row_vals):
                    continue
                rows.append(row_vals)
            if rows:
                df = pd.DataFrame(rows, columns=[cfg.COL_24Y_25N_COMPANY, cfg.COL_24Y_25N_2024, cfg.COL_24Y_25N_2025, cfg.COL_24Y_25N_2026, cfg.COL_24Y_25N_CARECRAFT, cfg.COL_24Y_25N_EMAIL_SENT, cfg.COL_24Y_25N_DATE])
                st.dataframe(df, use_container_width=True, height=400)
        except Exception as e:
            st.error(f"Failed to load 24Y/25N: {e}")

    with tab_labels:
        try:
            ws = wb[cfg.SHEET_LABEL_DEFINITION]
            rows = []
            for r in range(3, ws.max_row + 1):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, 5)]
                if all(v is None for v in row_vals):
                    continue
                rows.append(row_vals)
            if rows:
                df = pd.DataFrame(rows, columns=[cfg.COL_LABEL_DEF_IDX, cfg.COL_LABEL_DEF_LABEL, cfg.COL_LABEL_DEF_DEFINITION, cfg.COL_LABEL_DEF_NOTE])
                df = df.drop(columns=[cfg.COL_LABEL_DEF_IDX], errors="ignore")
                st.dataframe(df, use_container_width=True, hide_index=True)
        except Exception as e:
            st.error(f"Failed to load Labels: {e}")

    with tab_dupes:
        try:
            ws = wb[cfg.SHEET_DUPLICATES]
            rows = []
            for r in range(1, ws.max_row + 1):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, 3)]
                if all(v is None for v in row_vals):
                    continue
                rows.append(row_vals)
            if rows:
                df = pd.DataFrame(rows, columns=[cfg.COL_DUPLICATES_ACCOUNT, cfg.COL_DUPLICATES_NOTE])
                st.dataframe(df, use_container_width=True, hide_index=True)
        except Exception as e:
            st.error(f"Failed to load Duplicates: {e}")

elif page == "Merge":
    st.subheader("Merge ETL → Overview")
    st.markdown(
        "Upload your current `pivot_table.xlsx` + ETL output (CSV/XLSX) to merge. "
        "ETL quantities are added to the correct quarter column based on order date "
        "(Jan-Mar → Q1, Apr-Jun → Q2, Jul-Sep → Q3, Oct-Dec → Q4)."
    )

    pivot_file = st.file_uploader(
        "Pivot table file (xlsx)", type=["xlsx"], key="merge_pivot"
    )
    if pivot_file is None:
        st.info("Upload your pivot_table.xlsx first.")
        st.stop()

    import tempfile
    suffix = Path(pivot_file.name).suffix.lower()
    pivot_tmp = Path(f"/tmp/merge_pivot_{pivot_file.file_id}{suffix}")
    pivot_tmp.write_bytes(pivot_file.getvalue())

    # Point overview module at uploaded file (merge uses overview.PIVOT_PATH)
    import overview as overview_mod
    import merge as merge_mod
    overview_mod.PIVOT_PATH = str(pivot_tmp)

    if "etl_result" in st.session_state:
        st.info("Using ETL result from current session (sent from ETL tab).")
        etl = st.session_state["etl_result"]
        st.caption(f"ETL rows: {len(etl)}")
    else:
        merge_file = st.file_uploader(
            "ETL output file (csv / xlsx)", type=["csv", "xlsx"], key="merge_etl"
        )
        if merge_file is None:
            st.info("Send data from ETL tab via 'Send to Merge' button, or upload a file here.")
            st.stop()

        suffix = Path(merge_file.name).suffix.lower()
        tmp = Path(f"/tmp/merge_etl_{merge_file.file_id}{suffix}")
        tmp.write_bytes(merge_file.getvalue())

        try:
            if suffix == ".csv":
                etl = pd.read_csv(tmp)
            else:
                try:
                    etl = pd.read_excel(tmp, engine=cfg.EXCEL_ENGINES[0])
                except Exception:
                    etl = pd.read_excel(tmp, engine=cfg.EXCEL_ENGINES[1])
        except Exception as e:
            st.error(f"Failed to read file: {e}")
            st.stop()

    # ETL output may have renamed columns (e.g. "Sum of Lineitem quantity")
    if "Sum of Lineitem quantity" in etl.columns:
        qty_col = "Sum of Lineitem quantity"
    elif "Lineitem quantity" in etl.columns:
        qty_col = "Lineitem quantity"
    else:
        st.error("Missing column: Lineitem quantity")
        st.stop()
    required = {"Billing Company", qty_col, "Created at"}
    missing = required - set(etl.columns)
    if missing:
        st.error(f"Missing columns: {missing}")
        st.stop()

    # Detect new accounts in ETL not in Overview
    import overview as overview_mod
    existing_accounts = set(overview_mod.load_overview()["Account"])
    etl_accounts = set(etl["Billing Company"].dropna().unique())
    new_accounts = etl_accounts - existing_accounts
    if new_accounts:
        st.warning(f"**{len(new_accounts)} new account(s)** in ETL not in Overview: {', '.join(sorted(new_accounts)[:10])}{'...' if len(new_accounts) > 10 else ''}")

    from merge import compute_count, etl_date_label

    # Compute once — reused across tabs (deterministic for same ETL input)
    @st.cache_data
    def _compute_merged(etl_json: str):
        import json
        etl = pd.read_json(etl_json)
        return merge_overview_with_etl(etl).to_json()

    @st.cache_data
    def _compute_updated(etl_json: str):
        import json
        etl = pd.read_json(etl_json)
        return updated_overview_with_etl(etl).to_json()

    etl_json = etl.to_json()

    tab1, tab2, tab3 = st.tabs(["Side-by-side Merge", "Updated Overview", "Updated Count"])

    with tab1:
        merged = pd.read_json(_compute_merged(etl_json))
        new_count = merged[cfg.COL_QTY_2023].eq(0) & merged[cfg.COL_QTY_2024].eq(0) & merged[cfg.COL_QTY_2025].eq(0) & merged["ETL_Total"].gt(0)
        # Build dynamic ETL column list for display
        etl_cols = [c for c in ["ETL_Q1", "ETL_Q2", "ETL_Q3", "ETL_Q4", "ETL_Total"] if c in merged.columns]
        overview_cols = [c for c in merged.columns if c not in etl_cols]
        st.caption(f"Rows: {len(merged)} ({new_count.sum()} new) — ETL columns: {', '.join(etl_cols)}")
        st.dataframe(merged, use_container_width=True, height=600)

        csv_bytes = merged.to_csv(index=False).encode()
        st.download_button(
            "Download merged CSV", csv_bytes, "merged_overview.csv", "text/csv"
        )

    with tab2:
        updated = pd.read_json(_compute_updated(etl_json))
        new_count = updated[cfg.COL_QTY_2023].eq(0) & updated[cfg.COL_QTY_2024].eq(0) & updated[cfg.COL_QTY_2025].eq(0) & updated[cfg.COL_QTY_2026_TOTAL].gt(0)
        st.caption(f"Rows: {len(updated)} ({new_count.sum()} new) — ETL qty added to correct quarter, Qty_2026_ recomputed")
        st.dataframe(updated, use_container_width=True, height=600)

        csv_bytes = updated.to_csv(index=False).encode()
        st.download_button(
            "Download updated CSV", csv_bytes, "updated_overview.csv", "text/csv"
        )

        if st.button("Generate Updated Pivot XLSX"):
            import tempfile

            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                out_path = tmp.name
            try:
                merge_mod.generate_updated_pivot_xlsx(etl, out_path)
                with open(out_path, "rb") as f:
                    xlsx_bytes = f.read()
                st.download_button(
                    "Download Updated Pivot XLSX",
                    xlsx_bytes,
                    "Shopify_Pivot_Table_updated.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                st.success("Generated — click download above")
            except Exception as e:
                st.error(f"Failed to generate: {e}")

    with tab3:
        updated = pd.read_json(_compute_updated(etl_json))
        counts = compute_count(updated)
        label = etl_date_label(etl)

        st.caption(f"New column **{label}** appended to Count tab — categories computed from updated Overview")

        # Show as a simple table
        count_df = pd.DataFrame([
            {"Category": k, "Count": v}
            for k, v in counts.items()
        ])
        st.dataframe(count_df, use_container_width=True, hide_index=True)

        # Show delta vs last Count column if available
        try:
            from openpyxl import load_workbook as _lw
            wb_counts = _lw(pivot_tmp, data_only=True)
            ws_count = wb_counts["Count"]
            if ws_count.max_column >= 2:
                prev_col = ws_count.max_column  # last existing column (before our append)
                prev_label = ws_count.cell(row=1, column=prev_col).value
                prev_counts = {
                    ws_count.cell(row=r, column=1).value: ws_count.cell(row=r, column=prev_col).value
                    for r in range(2, 7)
                }
                delta_rows = []
                for cat, new_val in counts.items():
                    old_val = prev_counts.get(cat, 0) or 0
                    delta = new_val - old_val
                    delta_rows.append({
                        "Category": cat,
                        f"Before ({prev_label})": old_val,
                        f"After ({label})": new_val,
                        "Δ": delta,
                    })
                st.caption(f"Delta vs previous column ({prev_label})")
                st.dataframe(pd.DataFrame(delta_rows), use_container_width=True, hide_index=True)
        except Exception:
            pass
